"""Lee PDF/DOCX/PNG de la carpeta temarios/ y genera db.json."""

import json
import os
import platform
import sys
from pathlib import Path


def _python_venv() -> Path:
    raiz = Path(__file__).resolve().parent
    if platform.system().lower() == "windows":
        return raiz / ".venv" / "Scripts" / "python.exe"
    return raiz / ".venv" / "bin" / "python"


def _usar_venv_si_hace_falta() -> None:
    try:
        import fitz  # noqa: F401
        return
    except ImportError:
        pass

    venv_py = _python_venv()
    if venv_py.is_file() and venv_py.resolve() != Path(sys.executable).resolve():
        os.execv(str(venv_py), [str(venv_py), str(Path(__file__).resolve()), *sys.argv[1:]])

    print("Faltan dependencias (pymupdf, python-docx).")
    print("Ejecuta:  python run_app.py extract")
    print("O crea el venv:  python run_app.py")
    sys.exit(1)


_usar_venv_si_hace_falta()

import fitz
from docx import Document
from docx.oxml.ns import qn

CARPETA_DOCUMENTOS = Path(__file__).parent / "temarios"
ARCHIVO_DB = Path(__file__).parent / "db.json"
EXTENSIONES = {".pdf", ".docx", ".png"}
# Si una página tiene poco texto, se intenta OCR (PDF escaneado / diapositivas en imagen)
MIN_TEXTO_PAGINA = 40
MIN_TEXTO_CON_IMAGENES = 250
DPI_OCR = 300
MAX_LADO_PIXELES = 4500
OCR_MIN_CONFIANZA = 50
OCR_LADO_OBJETIVO = 2400  # diapositivas ~1376x768 → se escalan para Tesseract

_ocr_disponible: bool | None = None
_idiomas_ocr: str | None = None
_ocr_aviso_mostrado = False


def limpiar_texto(texto: str) -> str:
    return " ".join(texto.split())


def parrafos_desde_texto(texto: str) -> list[str]:
    texto = texto.replace("\r\n", "\n").replace("\r", "\n")
    parrafos: list[str] = []

    for bloque in texto.split("\n\n"):
        for linea in bloque.split("\n"):
            limpio = limpiar_texto(linea)
            if limpio:
                parrafos.append(limpio)

    if not parrafos and texto.strip():
        parrafos.append(limpiar_texto(texto))

    return parrafos


def ocr_disponible() -> bool:
    global _ocr_disponible
    if _ocr_disponible is not None:
        return _ocr_disponible
    try:
        import pytesseract

        pytesseract.get_tesseract_version()
        _ocr_disponible = True
    except Exception:
        _ocr_disponible = False
    return _ocr_disponible


def avisar_ocr_no_disponible() -> None:
    global _ocr_aviso_mostrado
    if _ocr_aviso_mostrado:
        return
    _ocr_aviso_mostrado = True
    print(
        "  [!] OCR no disponible. Para PDF escaneados instala Tesseract en el sistema:"
    )
    print("      Linux:  sudo apt install tesseract-ocr tesseract-ocr-spa")
    print("      Windows: https://github.com/UB-Mannheim/tesseract/wiki")


def idiomas_ocr() -> str:
    global _idiomas_ocr
    if _idiomas_ocr is not None:
        return _idiomas_ocr
    import pytesseract

    disponibles = set(pytesseract.get_languages(config=""))
    if "spa" in disponibles and "eng" in disponibles:
        _idiomas_ocr = "spa+eng"
    elif "spa" in disponibles:
        _idiomas_ocr = "spa"
    else:
        _idiomas_ocr = "eng"
    return _idiomas_ocr


def _matriz_render(dpi: int, ancho_pt: float, alto_pt: float) -> fitz.Matrix:
    zoom = dpi / 72.0
    if ancho_pt * zoom > MAX_LADO_PIXELES:
        zoom = MAX_LADO_PIXELES / ancho_pt
    if alto_pt * zoom > MAX_LADO_PIXELES:
        zoom = min(zoom, MAX_LADO_PIXELES / alto_pt)
    return fitz.Matrix(zoom, zoom)


def _pixmap_a_imagen(pixmap: fitz.Pixmap):
    from PIL import Image

    return Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)


def _luminosidad_media(imagen) -> float:
    from PIL import Image

    gris = imagen.convert("L")
    muestra = gris.resize(
        (min(160, gris.width), min(160, gris.height)), Image.Resampling.LANCZOS
    )
    pixels = muestra.get_flattened_data()
    return sum(pixels) / len(pixels)


def _preprocesar_imagen(imagen):
    from PIL import ImageEnhance, ImageOps

    if imagen.mode != "RGB":
        imagen = imagen.convert("RGB")

    gris = ImageOps.grayscale(imagen)
    media = _luminosidad_media(imagen)

    if media > 185:
        # Diapositivas UF1465: fondo blanco/claro, texto oscuro (sin binarizar)
        return ImageEnhance.Contrast(gris).enhance(1.25)
    if media < 95:
        gris = ImageOps.invert(gris)
        return ImageEnhance.Contrast(gris).enhance(1.4)

    gris = ImageOps.autocontrast(gris, cutoff=1)
    return ImageEnhance.Contrast(gris).enhance(1.35)


def _escalar_para_ocr(imagen):
    from PIL import Image

    ancho, alto = imagen.size
    lado = max(ancho, alto)
    if lado >= OCR_LADO_OBJETIVO:
        return imagen
    factor = min(2.5, OCR_LADO_OBJETIVO / lado)
    if factor <= 1.05:
        return imagen
    return imagen.resize(
        (int(ancho * factor), int(alto * factor)), Image.Resampling.LANCZOS
    )


def _es_linea_ocr_valida(linea: str) -> bool:
    linea = limpiar_texto(linea)
    if len(linea) < 5:
        return False
    alnum = sum(c.isalnum() for c in linea)
    if alnum < 4:
        return False
    ratio = alnum / len(linea)
    if ratio < 0.45:
        return False
    palabras_utiles = [
        p
        for p in linea.split()
        if len(p) >= 4 and any(c.isalpha() for c in p)
    ]
    if not palabras_utiles and len(linea) < 22:
        return False
    if len(linea) < 12 and ratio < 0.65:
        return False
    return True


def _ocr_lineas_desde_datos(datos: dict) -> dict[tuple, list[str]]:
    lineas: dict[tuple, list[str]] = {}
    for i, txt in enumerate(datos["text"]):
        txt = txt.strip()
        if not txt:
            continue
        try:
            conf = int(float(datos["conf"][i]))
        except (ValueError, TypeError):
            conf = -1
        if conf >= 0 and conf < OCR_MIN_CONFIANZA:
            continue
        clave = (
            datos["block_num"][i],
            datos["par_num"][i],
            datos["line_num"][i],
        )
        lineas.setdefault(clave, []).append(txt)
    return lineas


def _parrafos_desde_imagen(imagen) -> list[str]:
    import pytesseract

    imagen = _escalar_para_ocr(imagen.convert("RGB"))
    preparada = _preprocesar_imagen(imagen)
    lang = idiomas_ocr()
    vistos: set[str] = set()
    parrafos: list[str] = []

    def recoger(config: str) -> int:
        anadidos = 0
        try:
            datos = pytesseract.image_to_data(
                preparada,
                lang=lang,
                config=config,
                output_type=pytesseract.Output.DICT,
            )
        except pytesseract.TesseractError:
            datos = pytesseract.image_to_data(
                preparada,
                lang="eng",
                config=config,
                output_type=pytesseract.Output.DICT,
            )

        for palabras in _ocr_lineas_desde_datos(datos).values():
            linea = limpiar_texto(" ".join(palabras))
            if not _es_linea_ocr_valida(linea):
                continue
            clave = linea.lower()
            if clave in vistos:
                continue
            vistos.add(clave)
            parrafos.append(linea)
            anadidos += 1
        return anadidos

    base = "--oem 3 -c preserve_interword_spaces=1"
    recoger(f"{base} --psm 6")
    if len(parrafos) < 3:
        recoger(f"{base} --psm 11")

    return parrafos


def _imagenes_embebidas_pagina(pagina: fitz.Page) -> list:
    from PIL import Image
    import io

    documento = pagina.parent
    vistas: list = []
    vistos: set[int] = set()

    for info in pagina.get_images(full=True):
        xref = info[0]
        if xref in vistos:
            continue
        vistos.add(xref)
        try:
            datos = documento.extract_image(xref)
        except Exception:
            continue
        if datos.get("width", 0) < 80 or datos.get("height", 0) < 80:
            continue
        try:
            vistas.append(Image.open(io.BytesIO(datos["image"])))
        except Exception:
            continue

    return vistas


def _renderizar_pagina(pagina: fitz.Page):
    rect = pagina.rect
    matriz = _matriz_render(DPI_OCR, rect.width, rect.height)
    pixmap = pagina.get_pixmap(matrix=matriz, alpha=False)
    return _pixmap_a_imagen(pixmap)


def _es_diapositiva_solo_imagen(imagenes: list) -> bool:
    """PDF tipo PresentacionesUF1465: una PNG 16:9 por página, sin capa de texto."""
    if len(imagenes) != 1:
        return False
    ancho, alto = imagenes[0].size
    return ancho >= 600 and alto >= 400


def parrafos_desde_ocr(pagina: fitz.Page) -> list[str]:
    if not ocr_disponible():
        avisar_ocr_no_disponible()
        return []

    imagenes = _imagenes_embebidas_pagina(pagina)
    sin_texto_nativo = len(pagina.get_text().strip()) < MIN_TEXTO_PAGINA

    if imagenes and sin_texto_nativo and _es_diapositiva_solo_imagen(imagenes):
        return _parrafos_desde_imagen(imagenes[0])

    lineas_vistas: set[str] = set()
    parrafos: list[str] = []

    def anadir(nuevos: list[str]) -> None:
        for parrafo in nuevos:
            clave = parrafo.lower()
            if clave not in lineas_vistas and len(parrafo) >= 2:
                lineas_vistas.add(clave)
                parrafos.append(parrafo)

    for imagen in imagenes:
        anadir(_parrafos_desde_imagen(imagen))

    if sin_texto_nativo and not parrafos:
        anadir(_parrafos_desde_imagen(_renderizar_pagina(pagina)))

    return parrafos


def pagina_necesita_ocr(pagina: fitz.Page, parrafos: list[str]) -> bool:
    caracteres = sum(len(p) for p in parrafos)
    if caracteres < MIN_TEXTO_PAGINA:
        return True
    return bool(pagina.get_images()) and caracteres < MIN_TEXTO_CON_IMAGENES


def extraer_pdf(ruta: Path) -> list[dict]:
    registros: list[dict] = []
    documento = fitz.open(ruta)
    total_paginas = len(documento)
    paginas_ocr = 0

    for numero, pagina in enumerate(documento, start=1):
        parrafos = parrafos_desde_texto(pagina.get_text())

        if pagina_necesita_ocr(pagina, parrafos):
            ocr = parrafos_desde_ocr(pagina)
            if ocr:
                if parrafos:
                    existentes = {p.lower() for p in parrafos}
                    ocr = [p for p in ocr if p.lower() not in existentes]
                parrafos = parrafos + ocr
                paginas_ocr += 1
                print(f"  OCR página {numero}/{total_paginas}")

        for parrafo in parrafos:
            registros.append(
                {"archivo": ruta.name, "pagina": numero, "texto": parrafo}
            )

    documento.close()
    if paginas_ocr:
        print(f"  {paginas_ocr} página(s) con texto reconocido por OCR")
    return registros


def _tiene_salto_pagina(parrafo) -> bool:
    for run in parrafo.runs:
        for br in run._element.findall(
            ".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br"
        ):
            if br.get(qn("w:type")) == "page":
                return True
    return False


def extraer_docx(ruta: Path) -> list[dict]:
    registros: list[dict] = []
    documento = Document(ruta)
    pagina = 1

    for parrafo in documento.paragraphs:
        if _tiene_salto_pagina(parrafo):
            pagina += 1

        texto = limpiar_texto(parrafo.text)
        if texto:
            registros.append(
                {"archivo": ruta.name, "pagina": pagina, "texto": texto}
            )

    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [limpiar_texto(c.text) for c in fila.cells if c.text.strip()]
            if celdas:
                registros.append(
                    {
                        "archivo": ruta.name,
                        "pagina": pagina,
                        "texto": " | ".join(celdas),
                    }
                )

    return registros


def extraer_png(ruta: Path) -> list[dict]:
    from PIL import Image

    if not ocr_disponible():
        avisar_ocr_no_disponible()
        print(f"  [!] Sin OCR, se omite {ruta.name}")
        return []

    imagen = Image.open(ruta)
    parrafos = _parrafos_desde_imagen(imagen)
    if parrafos:
        print(f"  OCR: {len(parrafos)} línea(s) reconocidas")
    return [
        {"archivo": ruta.name, "pagina": 1, "texto": parrafo}
        for parrafo in parrafos
    ]


def extraer_archivo(ruta: Path) -> list[dict]:
    if ruta.suffix.lower() == ".pdf":
        return extraer_pdf(ruta)
    if ruta.suffix.lower() == ".docx":
        return extraer_docx(ruta)
    if ruta.suffix.lower() == ".png":
        return extraer_png(ruta)
    return []


def construir_base() -> list[dict]:
    if not CARPETA_DOCUMENTOS.is_dir():
        CARPETA_DOCUMENTOS.mkdir(parents=True)
        print(f"Carpeta creada: {CARPETA_DOCUMENTOS}")
        print("Coloca ahí tus PDF, DOCX y PNG y vuelve a ejecutar este script.")
        return []

    archivos = sorted(
        f for f in CARPETA_DOCUMENTOS.iterdir() if f.suffix.lower() in EXTENSIONES
    )

    if not archivos:
        print(f"No hay PDF, DOCX ni PNG en {CARPETA_DOCUMENTOS}")
        return []

    base: list[dict] = []
    for ruta in archivos:
        print(f"Procesando {ruta.name}...")
        base.extend(extraer_archivo(ruta))

    return base


def main() -> None:
    base = construir_base()
    with open(ARCHIVO_DB, "w", encoding="utf-8") as f:
        json.dump(base, f, ensure_ascii=False, indent=2)

    print(f"Listo: {len(base)} fragmentos guardados en {ARCHIVO_DB.name}")


if __name__ == "__main__":
    main()
